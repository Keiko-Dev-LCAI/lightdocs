"""Lightdocs backend — Notes → docs with multi-format output.

Flow: accept job → OCR → AIVM → parse optional ```lightdocs JSON → build
docx/md/xlsx/pptx → short-lived download. No secrets in responses.
Env: AIVM_RELAY, JOB_TTL_SECONDS, CORS_ORIGINS, LIGHTDOCS_DATA, PORT.
"""
from __future__ import annotations

import io
import json
import os
import re
import threading
import time
import uuid
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Optional

import requests
from docx import Document
from docx.shared import Pt
from flask import Flask, jsonify, request, send_file
from flask_cors import CORS

APP_NAME = "lightdocs"
VERSION = "0.6.0"

VALID_MODES = frozenset(
    {
        "notes-word",
        "notes-to-word",
        "meeting",
        "resume",
        "invoice",
        "sop",
        "study",
        "sheet",
        "deck",
        "explain",
        "rewrite",
        "condition",
        "dao",
        "litepaper",
        "announce",
    }
)

AIVM_RELAY = os.environ.get(
    "AIVM_RELAY", "https://web-production-aaaba.up.railway.app"
).rstrip("/")
JOB_TTL_SECONDS = int(os.environ.get("JOB_TTL_SECONDS", "3600"))
DATA_DIR = Path(os.environ.get("LIGHTDOCS_DATA", "/tmp/lightdocs-jobs"))
DATA_DIR.mkdir(parents=True, exist_ok=True)
_META_DIR = DATA_DIR / "meta"
_META_DIR.mkdir(parents=True, exist_ok=True)

# Monthly subscription ($1/mo) + one-time free docs per wallet — not credits
SUB_PRICE_USD = float(os.environ.get("LIGHTDOCS_SUB_PRICE_USD") or "1.00")
KEIKO_DISCOUNT = float(os.environ.get("LIGHTDOCS_KEIKO_DISCOUNT") or "0.20")
FREE_DOCS_TOTAL = int(os.environ.get("LIGHTDOCS_FREE_DOCS_TOTAL") or "100")
PASS_DAYS = int(os.environ.get("LIGHTDOCS_PASS_DAYS") or "30")
KEIKO_RECEIVE_WALLET = (os.environ.get("KEIKO_RECEIVE_WALLET") or "").strip()
LCAI_RECEIVE_WALLET = (os.environ.get("LCAI_RECEIVE_WALLET") or "").strip()
_PASS_DIR = DATA_DIR / "passes"
_PASS_DIR.mkdir(parents=True, exist_ok=True)
_rate_lock = threading.Lock()
_rate_hits: dict[str, list[float]] = {}
_lcai_price_cache = {"price": 0.004, "ts": 0.0}
_keiko_price_cache = {"price": 0.00001, "ts": 0.0}

CORS_ORIGINS = [
    o.strip()
    for o in os.environ.get(
        "CORS_ORIGINS",
        "https://keiko-dev-lcai.github.io,http://localhost:5500,http://127.0.0.1:5500,http://localhost:8080",
    ).split(",")
    if o.strip()
]

app = Flask(__name__)
CORS(app, origins=CORS_ORIGINS, supports_credentials=False)

_jobs_lock = threading.Lock()
_jobs: dict[str, dict[str, Any]] = {}


def _job_meta_path(job_id: str) -> Path:
    return _META_DIR / f"{job_id}.json"


def _save_job(job_id: str, job: dict[str, Any]) -> None:
    try:
        _job_meta_path(job_id).write_text(json.dumps(job), encoding="utf-8")
    except Exception as e:
        print(f"[job] meta save failed: {e}")


def _load_job(job_id: str) -> Optional[dict[str, Any]]:
    with _jobs_lock:
        if job_id in _jobs:
            return _jobs[job_id]
    p = _job_meta_path(job_id)
    if not p.is_file():
        return None
    try:
        job = json.loads(p.read_text(encoding="utf-8"))
        with _jobs_lock:
            _jobs[job_id] = job
        return job
    except Exception:
        return None


def _unlink_quiet(path: Path) -> None:
    try:
        path.unlink(missing_ok=True)
    except OSError:
        pass


def _cleanup_expired() -> None:
    now = time.time()
    with _jobs_lock:
        dead = [
            jid
            for jid, j in _jobs.items()
            if now - float(j.get("created", now)) > JOB_TTL_SECONDS
        ]
        for jid in dead:
            for key in ("docx_path", "file_path"):
                path = _jobs[jid].get(key)
                if path and Path(path).is_file():
                    _unlink_quiet(Path(path))
            _unlink_quiet(_job_meta_path(jid))
            _jobs.pop(jid, None)
    for p in _META_DIR.glob("*.json"):
        try:
            job = json.loads(p.read_text(encoding="utf-8"))
            if now - float(job.get("created", 0)) > JOB_TTL_SECONDS:
                for key in ("docx_path", "file_path"):
                    doc = job.get(key)
                    if doc and Path(doc).is_file():
                        _unlink_quiet(Path(doc))
                _unlink_quiet(p)
        except Exception:
            pass


def ocr_image_bytes(data: bytes) -> str:
    try:
        from rapidocr_onnxruntime import RapidOCR
        from PIL import Image

        img = Image.open(io.BytesIO(data)).convert("RGB")
        buf = io.BytesIO()
        img.save(buf, format="PNG")
        ocr = RapidOCR()
        result, _ = ocr(buf.getvalue())
        if not result:
            return ""
        return "\n".join(line[1] for line in result if line and len(line) > 1).strip()
    except Exception as e:
        print(f"[ocr] failed: {e}")
        return ""


def aivm_infer(prompt: str, timeout: int = 240) -> str:
    start = requests.post(
        f"{AIVM_RELAY}/api/chat",
        json={"message": prompt, "mode": "chat"},
        timeout=30,
    )
    if not start.ok:
        raise RuntimeError(f"AIVM start failed: {start.status_code} {start.text[:200]}")
    data = start.json()
    job_id = data.get("job_id")
    if not job_id:
        return (data.get("reply") or data.get("message") or "").strip()
    deadline = time.time() + timeout
    while time.time() < deadline:
        time.sleep(5)
        poll = requests.get(
            f"{AIVM_RELAY}/api/chat/status",
            params={"job_id": job_id},
            timeout=20,
        )
        if not poll.ok:
            raise RuntimeError(f"AIVM poll failed: {poll.status_code}")
        pd = poll.json()
        if pd.get("status") == "done":
            return (pd.get("reply") or "").strip()
        if pd.get("status") == "error":
            raise RuntimeError(pd.get("error") or "AIVM job failed")
    raise RuntimeError("AIVM timed out — try again")


STYLE_PROMPTS = {
    "formal": "Write in a formal, professional tone.",
    "plain": "Write in clear plain English for everyday readers.",
    "simple": "Write simply (ELI5) — short sentences, no jargon.",
    "nerd": "Write in a detailed, nerdy, precise tone.",
    "tech": "Write as a technical expert — precise terminology OK.",
    "friendly": "Write in a friendly, casual tone.",
}


_TELEMETRY_RE = re.compile(
    r"\{[^{}]*\"(?:promptTokens|evalTokens|tokensPerSecond|totalMs)\"[^{}]*\}"
)


_LEAK_LINE_RES = [
    re.compile(p, re.I)
    for p in (
        r"^here is the output",
        r"^here's the output",
        r"^i hope this helps",
        r"^let me know if",
        r"^sure,",
        r"^of course,",
        r"^turn messy notes into a clean document",
        r"^chart requested by user",
        r"^use only the input",
        r"^output shaping\s*:?\s*$",
        r"^do not invent",
        r"^critical\s*:",
        r"^critical facts\s*:?\s*$",
        r"^task summary\s*:?\s*$",
        r"^document content\s*:?\s*$",
        r"^numeric series\b.*requested by user",
        r"^mode\s*:\s*\w+",
        r"^task\s*:?\s*$",
        r"^you are lightdocs",
        r"^format only the text inside",
        r"^never restate",
        r"^no preamble",
        r"^supported chart types",
        r"^if input has no clear",
        r"^pick the best of bar",
        r"^proposal type hint",
    )
]


def _strip_aivm_noise(text: str) -> str:
    """Drop AIVM chatter/telemetry and leaked prompt/instruction lines."""
    t = _TELEMETRY_RE.sub("", text or "")
    lines = []
    for ln in t.splitlines():
        stripped = ln.strip()
        if not stripped:
            lines.append(ln)
            continue
        low = stripped.lower()
        if any(rx.search(stripped) for rx in _LEAK_LINE_RES):
            continue
        # bare section headers that mirror the prompt scaffold
        if low in (
            "task",
            "task:",
            "critical",
            "critical:",
            "input",
            "input:",
            "output shaping",
            "output shaping:",
            "mode",
            "mode:",
        ):
            continue
        lines.append(ln)
    # collapse excess blank lines
    out: list[str] = []
    blank = 0
    for ln in lines:
        if not ln.strip():
            blank += 1
            if blank <= 2:
                out.append(ln)
        else:
            blank = 0
            out.append(ln)
    return "\n".join(out).strip()


def try_parse_numeric_series(raw: str) -> Optional[dict[str, Any]]:
    """If INPUT is a clear bare numeric series, return labels/values (never invent)."""
    text = (raw or "").strip()
    if not text:
        return None
    # reject long prose
    if len(text) > 400 or text.count("\n") > 40:
        return None
    words = re.findall(r"[A-Za-z]{4,}", text)
    if len(words) > 8:
        return None
    # labeled pairs: North 120 / North: 120 / North,120
    pairs = re.findall(
        r"(?m)^\s*([A-Za-z][A-Za-z0-9 _/-]{0,24}?)\s*[:=\-,]?\s*(-?\d+(?:\.\d+)?)\s*$",
        text,
    )
    labels: list[str] = []
    values: list[float] = []
    if len(pairs) >= 2:
        for lab, val in pairs:
            labels.append(lab.strip())
            values.append(float(val))
    else:
        # bare numbers separated by comma/space/newline
        nums = re.findall(r"(?<![A-Za-z])-?\d+(?:\.\d+)?(?![A-Za-z])", text)
        # require that almost all tokens are numbers
        tokens = [t for t in re.split(r"[\s,;]+", text) if t]
        if len(nums) < 2:
            return None
        if len(tokens) and len(nums) < max(2, int(len(tokens) * 0.6)):
            return None
        values = [float(n) for n in nums[:24]]
        labels = [f"Item {i+1}" for i in range(len(values))]
    if len(labels) != len(values) or len(values) < 2:
        return None
    return {
        "title": "Values",
        "labels": labels,
        "values": values,
    }


def _simple_doc_from_input(raw: str) -> str:
    """Short titled markdown from INPUT only (used when AIVM echoes instructions)."""
    text = (raw or "").strip()
    if not text:
        return "# Notes\n\n"
    lines = [ln.strip() for ln in text.splitlines() if ln.strip()]
    nums = try_parse_numeric_series(text)
    if nums:
        body = "\n".join(
            f"- {lab}: {val:g}" for lab, val in zip(nums["labels"], nums["values"])
        )
        return f"# Values\n\n{body}\n"
    if len(lines) <= 12:
        bullets = "\n".join(f"- {ln}" for ln in lines)
        return f"# Notes\n\n{bullets}\n"
    return f"# Notes\n\n{text[:4000]}\n"


def _looks_like_instruction_leak(md: str) -> bool:
    low = (md or "").lower()
    hits = 0
    for needle in (
        "turn messy notes into a clean document",
        "chart requested by user",
        "output shaping",
        "critical facts",
        "task summary",
        "numeric series",
        "use only the input",
        "you are lightdocs",
    ):
        if needle in low:
            hits += 1
    return hits >= 1


def parse_lightdocs_payload(raw: str) -> tuple[str, dict[str, Any]]:
    """Split AIVM output into markdown + optional ```lightdocs JSON block."""
    meta: dict[str, Any] = {}
    md = raw
    m = re.search(r"```lightdocs\s*([\s\S]*?)```", raw, re.I)
    if m:
        md = (raw[: m.start()] + raw[m.end() :]).strip()
        try:
            meta = json.loads(m.group(1).strip())
            if not isinstance(meta, dict):
                meta = {}
        except Exception:
            meta = {}
    # also accept bare JSON object with charts/sheets/slides anywhere
    if not meta:
        m2 = re.search(
            r"(\{\s*\"(?:charts|sheets|slides)\"\s*:\s*\[[\s\S]*?\]\s*\})", raw
        )
        if m2:
            try:
                meta = json.loads(m2.group(1))
                md = (raw[: m2.start()] + raw[m2.end() :]).strip()
            except Exception:
                pass
    md = _strip_aivm_noise(md)
    return md, meta


MODE_TASKS: dict[str, str] = {
    "notes-word": (
        "Turn messy notes into a clean document: short H1 title, headings, bullets, "
        "and short paragraphs. Preserve facts from INPUT only."
    ),
    "notes-to-word": (
        "Turn messy notes into a clean document: short H1 title, headings, bullets, "
        "and short paragraphs. Preserve facts from INPUT only."
    ),
    "meeting": (
        "Turn messy meeting notes into: (1) Meeting title + meta if present, "
        "(2) Concise summary, (3) Decisions, (4) Action items as a checklist with "
        "Owner — Task — Due when present (use [TO FILL] if owner/due missing). "
        "Do not invent attendees or commitments."
    ),
    "resume": (
        "Build a structured resume/CV OR a cover letter from INPUT. "
        "If INPUT looks like a role + notes, prefer a cover letter; otherwise a resume "
        "with Contact, Summary, Experience, Education, Skills. Use [TO FILL] for missing "
        "contact details — never invent employers, dates, or credentials."
    ),
    "invoice": (
        "Draft an invoice or quote: seller/buyer placeholders, invoice/quote number "
        "[TO FILL], date, line items (description, qty, unit price, line total), "
        "subtotal, tax if mentioned, total. Currency from INPUT or mark [TO FILL]. "
        "Do not invent prices."
    ),
    "sop": (
        "Turn rough steps into a clear SOP or checklist: purpose, prerequisites, "
        "numbered procedure steps, and a final checkbox verification list. "
        "Keep steps actionable; mark unknowns as [TO FILL]."
    ),
    "study": (
        "Build a study guide / lesson outline from INPUT: learning objectives, "
        "outline of topics, key terms with short definitions, and review questions. "
        "Stay on the subject in INPUT."
    ),
    "sheet": (
        "Shape INPUT as spreadsheet data: infer sensible columns and rows from the notes. "
        "Prefer real numbers/labels from INPUT; do not invent metrics."
    ),
    "deck": (
        "Turn INPUT into a short presentation outline: 5–10 slides with titles and "
        "tight bullets. One idea per slide; preserve facts from INPUT only."
    ),
    "explain": (
        "Scan & explain: give a plain-language explanation of what the INPUT document "
        "or notes mean. Structure as Overview, Key points, What to do next. "
        "If quoting, keep short and clearly marked. Do not invent content not in INPUT."
    ),
    "rewrite": (
        "Rewrite / simplify / translate the INPUT text in the requested voice. "
        "Preserve meaning and facts; improve clarity. If INPUT asks for a language, "
        "translate into that language; otherwise rewrite in the voice style."
    ),
    "condition": (
        "Create a property condition / move-in report from notes/photo OCR: "
        "Property/unit [TO FILL], date, rooms/areas with condition notes, issues found, "
        "photos referenced by filename if present, and a signature/ack placeholder. "
        "Do not invent damage that is not in INPUT."
    ),
    "dao": (
        "Draft a Lightchain LCAI DUNA governor proposal from rough notes. Match real "
        "Lightchain proposal structure EXACTLY with these sections in order:\n"
        "1) Title — [Verb] [Subject], plain and specific (no hype)\n"
        "2) Summary — 1–2 sentences of what is authorized, including headline numbers\n"
        "3) Key terms — bullets for amounts, rates, counterparties, duration\n"
        "4) Why it matters / rationale — short paragraph\n"
        "5) Scope & limitations — what this does NOT authorize\n"
        "6) What this means / effect if passed — bullets\n"
        "7) Governance / execution footer — voting-window note + placeholders for "
        "on-chain record and dao.lightchain.ai link\n"
        "House style: sober, transparent, DUNA framing (LCAI DUNA; Administrator "
        "Quantum Counsel LLC when relevant). NEVER invent tx hashes, proposal IDs, "
        "treasury balances, or vote results — use [TO FILL]."
    ),
    "litepaper": (
        "Write a Lightchain-oriented litepaper / one-pager from INPUT: Title, Hook, "
        "Problem, Solution, How it works, Token/utility if mentioned, Roadmap, "
        "Call to action. Keep it one-page dense. No invented tokenomics numbers."
    ),
    "announce": (
        "Write a Forum / Discord announcement in Lightchain community style: short, "
        "scannable, with emoji section markers (e.g. 🚀 🚨 🧠 ⏳ 🗳️) where natural, "
        "clear CTA, and link placeholders as [TO FILL]. Do not invent proposal IDs "
        "or tx hashes."
    ),
}


def _format_instructions(output: str, mode: str) -> str:
    """How the model should shape the reply for the chosen download format."""
    if output == "xlsx":
        return (
            "Return ONLY a fenced block (nothing else):\n"
            "```lightdocs\n"
            '{ "sheets": [ { "name": "Sheet1", "columns": ["A","B"], '
            '"rows": [["x",1],["y",2]] } ] }\n'
            "```\n"
            "If INPUT is not tabular, one Notes column listing the points."
        )
    if output == "pptx":
        return (
            "Return ONLY a fenced block (nothing else):\n"
            "```lightdocs\n"
            '{ "slides": [ { "title": "Overview", "bullets": ["Point A","Point B"], '
            '"notes": "" } ] }\n'
            "```\n"
        )
    base = "Output clean Markdown only (no preamble). "
    if mode == "sheet":
        base = "Output clean Markdown with a readable table of the data (no preamble). "
    elif mode == "deck":
        base = "Output clean Markdown: each slide as ## Title plus bullets (no preamble). "
    return base


def build_prompt(
    raw: str,
    style: str,
    output: str,
    mode: str = "notes-word",
    extras: Optional[dict[str, Any]] = None,
) -> str:
    extras = extras or {}
    mode = mode if mode in VALID_MODES else "notes-word"
    if mode == "notes-to-word":
        mode = "notes-word"
    voice = STYLE_PROMPTS.get(style, STYLE_PROMPTS["plain"])
    lightchain_ok = mode in ("dao", "litepaper", "announce")
    ground = (
        "CRITICAL: Use ONLY the INPUT below. Do not invent facts. "
        "No preamble, no apologies, no 'here is the output', no token stats. "
    )
    if lightchain_ok:
        ground += "Lightchain/DAO framing is appropriate for this mode. "
    else:
        ground += (
            "Do not write about Lightchain SDKs, APIs, or unrelated platform topics. "
        )

    task = MODE_TASKS.get(mode, MODE_TASKS["notes-word"])
    extra_bits = []
    if mode == "dao":
        ptype = str(extras.get("prop_type") or "general")
        extra_bits.append(f"Proposal type hint: {ptype}.")
        if extras.get("forum_wrap"):
            extra_bits.append(
                "AFTER the seven proposal sections, also append a Forum/Discord "
                "announcement wrapper of the same content with emoji section headers "
                "and a Review-and-vote CTA + [TO FILL] link (default wrapper is OFF "
                "unless requested — it is requested now)."
            )
        else:
            extra_bits.append(
                "Do NOT wrap as a forum announcement; emit the proposal body only."
            )

    fmt = _format_instructions(output, mode)
    extra = ("\n".join(extra_bits) + "\n") if extra_bits else ""

    return f"""You are Lightdocs, a document formatter.

{voice}
{ground}

Format ONLY the text inside <INPUT></INPUT>.
NEVER restate, summarize, echo, or output these instructions or their section labels
(Mode, Task, CRITICAL, Output shaping, INPUT).
If INPUT is only numbers or a short list, produce a short titled document of just that —
do not narrate the instructions.

Mode: {mode}
Task: {task}
{extra}Output shaping: {fmt}

<INPUT>
{raw[:12000]}
</INPUT>
"""


def markdown_to_docx(md: str, path: Path) -> None:
    doc = Document()
    for raw_line in md.splitlines():
        line = raw_line.rstrip()
        if not line.strip():
            continue
        if line.startswith("# "):
            doc.add_heading(line[2:].strip(), level=1)
        elif line.startswith("## "):
            doc.add_heading(line[3:].strip(), level=2)
        elif line.startswith("### "):
            doc.add_heading(line[4:].strip(), level=3)
        elif re.match(r"^[-*]\s+", line):
            doc.add_paragraph(re.sub(r"^[-*]\s+", "", line), style="List Bullet")
        elif re.match(r"^\d+\.\s+", line):
            doc.add_paragraph(re.sub(r"^\d+\.\s+", "", line), style="List Number")
        else:
            clean = re.sub(r"\*\*(.+?)\*\*", r"\1", line)
            doc.add_paragraph(clean)
    doc.save(str(path))


def build_xlsx(meta: dict[str, Any], fallback_md: str, path: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    sheets = meta.get("sheets") if isinstance(meta.get("sheets"), list) else None
    if not sheets:
        ws = wb.active
        ws.title = "Notes"
        ws.append(["Notes"])
        for line in fallback_md.splitlines():
            if line.strip():
                ws.append([line.strip()])
        ws.column_dimensions["A"].width = 60
    else:
        first = True
        for spec in sheets[:8]:
            name = str(spec.get("name") or "Sheet")[:31] or "Sheet"
            cols = list(spec.get("columns") or [])
            rows = list(spec.get("rows") or [])
            if first:
                ws = wb.active
                ws.title = name
                first = False
            else:
                ws = wb.create_sheet(name)
            if cols:
                ws.append(cols)
                for cell in ws[1]:
                    cell.font = Font(bold=True, color="5B4BFF")
            for row in rows:
                if isinstance(row, (list, tuple)):
                    ws.append(list(row))
                else:
                    ws.append([row])
            for i, _ in enumerate(cols or [0], start=1):
                ws.column_dimensions[chr(64 + min(i, 26))].width = 16
    wb.save(str(path))


def build_pptx(meta: dict[str, Any], fallback_md: str, path: Path) -> None:
    from pptx import Presentation
    from pptx.dml.color import RGBColor as PptRGB
    from pptx.util import Pt as PptPt

    prs = Presentation()
    slides = meta.get("slides") if isinstance(meta.get("slides"), list) else None
    if not slides:
        # fallback: split markdown headings into slides
        chunks = re.split(r"\n(?=# )", fallback_md.strip()) or [fallback_md]
        slides = []
        for ch in chunks[:12]:
            lines = [ln.strip() for ln in ch.splitlines() if ln.strip()]
            title = lines[0].lstrip("# ").strip() if lines else "Notes"
            bullets = [re.sub(r"^[-*#\d.\s]+", "", ln) for ln in lines[1:8]]
            slides.append({"title": title, "bullets": [b for b in bullets if b], "notes": ""})
    for spec in slides[:16]:
        layout = prs.slide_layouts[1]  # title + content
        slide = prs.slides.add_slide(layout)
        title = str(spec.get("title") or "Slide")
        slide.shapes.title.text = title
        body = slide.placeholders[1].text_frame
        body.clear()
        bullets = list(spec.get("bullets") or [])
        if not bullets:
            p = body.paragraphs[0]
            p.text = ""
        else:
            for i, b in enumerate(bullets[:10]):
                p = body.paragraphs[0] if i == 0 else body.add_paragraph()
                p.text = str(b)
                p.level = 0
                p.font.size = PptPt(20)
                p.font.color.rgb = PptRGB(0x14, 0x15, 0x2C)
        notes = str(spec.get("notes") or "").strip()
        if notes:
            slide.notes_slide.notes_text_frame.text = notes
    prs.save(str(path))


def process_job(
    job_id: str,
    text: str,
    images: list[bytes],
    style: str,
    output: str,
    mode: str = "notes-word",
    extras: Optional[dict[str, Any]] = None,
) -> None:
    def set_step(status: str, step: str, **extra: Any) -> None:
        with _jobs_lock:
            job = _jobs.get(job_id) or {}
            job.update({"status": status, "step": step, **extra})
            _jobs[job_id] = job
            _save_job(job_id, job)

    extras = extras or {}
    try:
        set_step("ocr", "Reading images (OCR)…")
        ocr_parts = []
        for i, blob in enumerate(images[:5]):
            t = ocr_image_bytes(blob)
            if t:
                ocr_parts.append(t)
            set_step("ocr", f"OCR image {i + 1}/{min(len(images), 5)}…")

        combined = (text or "").strip()
        if ocr_parts:
            combined = (combined + "\n\n" + "\n\n".join(ocr_parts)).strip()
        if not combined:
            raise RuntimeError("No text found — paste notes or use a clearer photo.")

        set_step("aivm", f"AIVM drafting ({mode})…")
        raw_out = aivm_infer(build_prompt(combined, style, output, mode, extras))
        if not raw_out or len(raw_out) < 4:
            raise RuntimeError("AIVM returned empty output — try again.")
        md, meta = parse_lightdocs_payload(raw_out)
        if not md.strip() and output in ("docx", "md"):
            md = _strip_aivm_noise(raw_out)
        if _looks_like_instruction_leak(md) or not md.strip():
            md = _simple_doc_from_input(combined)

        # Bare numeric input: never trust model prose (it hallucinates on trivial input).
        # Build a clean Values list from the user's own numbers — no charts.
        if try_parse_numeric_series(combined):
            md = _simple_doc_from_input(combined)
            meta.pop("charts", None)

        # mode defaults: sheet→xlsx builder path if sheets present even when md empty
        if mode == "sheet" and output not in ("xlsx", "pptx") and meta.get("sheets"):
            # keep user's output; markdown_to_docx/md will use md or a simple dump
            if not md.strip():
                md = json.dumps(meta.get("sheets"), indent=2)
        if mode == "deck" and output not in ("pptx",) and meta.get("slides") and not md.strip():
            md = json.dumps(meta.get("slides"), indent=2)

        set_step("building", f"Building {output} file…")
        if output == "md":
            fname = f"lightdocs-{job_id[:8]}.md"
            path = DATA_DIR / fname
            path.write_text(md.strip() + "\n", encoding="utf-8")
            mime = "text/markdown; charset=utf-8"
            text_out = md.strip()
        elif output == "xlsx":
            fname = f"lightdocs-{job_id[:8]}.xlsx"
            path = DATA_DIR / fname
            build_xlsx(meta, md, path)
            mime = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            text_out = (
                json.dumps(meta.get("sheets") or [], indent=2)
                if meta.get("sheets")
                else md.strip()
            )
        elif output == "pptx":
            fname = f"lightdocs-{job_id[:8]}.pptx"
            path = DATA_DIR / fname
            build_pptx(meta, md, path)
            mime = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
            text_out = (
                json.dumps(meta.get("slides") or [], indent=2)
                if meta.get("slides")
                else md.strip()
            )
        else:
            output = "docx"
            fname = f"lightdocs-{job_id[:8]}.docx"
            path = DATA_DIR / fname
            markdown_to_docx(md, path)
            mime = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
            text_out = md.strip()

        set_step(
            "done",
            "Done — download ready. We don’t keep your docs long.",
            text_out=text_out,
            download_name=fname,
            file_path=str(path),
            docx_path=str(path),  # backward compat
            mime=mime,
            output=output,
            error=None,
        )
    except Exception as e:
        print(f"[job {job_id}] error: {e}")
        set_step("error", "Failed", error=str(e)[:400])


@app.get("/api/health")
def health():
    return jsonify(
        {
            "status": "ok",
            "app": APP_NAME,
            "version": VERSION,
            "aivm_relay_configured": bool(AIVM_RELAY),
            "retention_seconds": JOB_TTL_SECONDS,
            "formats": ["docx", "md", "xlsx", "pptx"],
            "modes": sorted(m for m in VALID_MODES if m != "notes-to-word"),
            "pay_model": "subscription",
            "free_docs_total": FREE_DOCS_TOTAL,
            "sub_price_usd": SUB_PRICE_USD,
            "privacy": "Uploads processed for the job only; files auto-expire. We don’t keep your docs. Drafts stay on your device only. Wallet login stays on-device.",
        }
    )


def _truthy(val: Any) -> bool:
    if isinstance(val, bool):
        return val
    return str(val or "").strip().lower() in ("1", "true", "yes", "on")


@app.post("/api/jobs")
def create_job():
    _cleanup_expired()
    style = "plain"
    text = ""
    output = "docx"
    mode = "notes-word"
    prop_type = "general"
    forum_wrap = False
    device_id = (request.headers.get("X-Device-Id") or "").strip()
    wallet = (request.headers.get("X-Wallet") or "").strip().lower()
    images: list[bytes] = []

    if request.content_type and "multipart/form-data" in request.content_type:
        text = (request.form.get("text") or "").strip()
        style = (request.form.get("style") or "plain").strip()
        output = (request.form.get("output") or request.form.get("outfmt") or "docx").strip()
        mode = (request.form.get("mode") or "notes-word").strip()
        prop_type = (request.form.get("prop_type") or request.form.get("propType") or "general").strip()
        forum_wrap = _truthy(request.form.get("forum_wrap") or request.form.get("forumWrap"))
        device_id = (request.form.get("device_id") or device_id).strip()
        wallet = (request.form.get("wallet") or request.form.get("walletAddress") or wallet).strip().lower()
        for key in ("images", "image", "files"):
            for f in request.files.getlist(key):
                if f and f.filename:
                    data = f.read()
                    if data and len(data) < 12_000_000:
                        images.append(data)
    else:
        body = request.get_json(silent=True) or {}
        text = (body.get("text") or "").strip()
        style = (body.get("style") or "plain").strip()
        output = (body.get("output") or body.get("outfmt") or "docx").strip()
        mode = (body.get("mode") or "notes-word").strip()
        prop_type = str(body.get("prop_type") or body.get("propType") or "general").strip()
        forum_wrap = _truthy(body.get("forum_wrap") or body.get("forumWrap"))
        device_id = str(body.get("device_id") or device_id).strip()
        wallet = str(body.get("wallet") or body.get("walletAddress") or wallet).strip().lower()
        for b64 in body.get("images") or []:
            try:
                import base64

                raw = b64.split(",", 1)[-1]
                images.append(base64.b64decode(raw))
            except Exception:
                pass

    if output not in ("docx", "md", "xlsx", "pptx"):
        output = "docx"
    if mode not in VALID_MODES:
        mode = "notes-word"
    if mode == "notes-to-word":
        mode = "notes-word"
    if prop_type not in ("general", "treasury", "param", "signal", "grant"):
        prop_type = "general"

    if not text and not images:
        return jsonify({"error": "Provide text and/or at least one image"}), 400

    device_id = str(device_id or "").strip() or "anon"
    ok_pay, pay_msg, pay_meta = _authorize_job(wallet, device_id=device_id)
    if not ok_pay:
        return jsonify({"error": pay_msg, "need_pay": bool(pay_meta.get("need_pay")), **pay_meta}), 402

    extras = {
        "prop_type": prop_type,
        "forum_wrap": forum_wrap,
    }
    job_id = uuid.uuid4().hex
    job = {
        "id": job_id,
        "status": "queued",
        "step": "Queued",
        "created": time.time(),
        "error": None,
        "text_out": None,
        "download_name": None,
        "file_path": None,
        "docx_path": None,
        "mime": None,
        "mode": mode,
        "style": style,
        "output": output,
        "extras": extras,
        "pay": {"path": pay_meta.get("path"), "paid": bool(pay_meta.get("paid"))},
    }
    with _jobs_lock:
        _jobs[job_id] = job
        _save_job(job_id, job)

    threading.Thread(
        target=process_job,
        args=(job_id, text, images, style, output, mode, extras),
        daemon=True,
    ).start()
    return jsonify(
        {
            "job_id": job_id,
            "status": "queued",
            "poll_url": f"/api/jobs/{job_id}",
            "retention_seconds": JOB_TTL_SECONDS,
            "output": output,
            "mode": mode,
            "pay": pay_meta,
            "pay_status": pay_msg,
        }
    ), 202


@app.get("/api/jobs/<job_id>")
def get_job(job_id: str):
    _cleanup_expired()
    job = _load_job(job_id)
    if not job:
        return jsonify({"error": "Job not found or expired"}), 404
    return jsonify(
        {
            "job_id": job_id,
            "status": job["status"],
            "step": job["step"],
            "error": job.get("error"),
            "text": job.get("text_out") if job["status"] == "done" else None,
            "download_url": (
                f"/api/jobs/{job_id}/download" if job["status"] == "done" else None
            ),
            "download_name": job.get("download_name"),
            "output": job.get("output"),
            "mode": job.get("mode"),
            "privacy": "We don’t keep your docs — downloads expire automatically. Drafts stay on your device only.",
        }
    )


@app.get("/api/jobs/<job_id>/download")
def download_job(job_id: str):
    _cleanup_expired()
    job = _load_job(job_id)
    if not job or job["status"] != "done":
        return jsonify({"error": "Not ready or expired"}), 404
    path = job.get("file_path") or job.get("docx_path")
    if not path or not Path(path).is_file():
        return jsonify({"error": "File gone — please generate again"}), 410
    return send_file(
        path,
        as_attachment=True,
        download_name=job.get("download_name") or "lightdocs-out.bin",
        mimetype=job.get("mime") or "application/octet-stream",
    )


def _decode_image_src(src: str) -> Optional[bytes]:
    if not src or not isinstance(src, str):
        return None
    try:
        if src.startswith("data:"):
            import base64

            raw = src.split(",", 1)[-1]
            return base64.b64decode(raw)
        if src.startswith("http://") or src.startswith("https://"):
            r = requests.get(src, timeout=30)
            if r.ok:
                return r.content
        # relative API path
        if src.startswith("/"):
            # not fetchable from here without host — skip
            return None
    except Exception as e:
        print(f"[export] image decode failed: {e}")
    return None


def _add_picture_wrapped(doc: Document, img_bytes: bytes, width_in: float, wrap: str) -> None:
    """Embed PNG/JPEG with approximate wrap: full/inline block, left/right via 2-col table."""
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches as DocInches

    width_in = max(0.8, min(float(width_in or 4.5), 6.5))
    wrap = (wrap or "full").lower()
    bio = io.BytesIO(img_bytes)
    if wrap in ("left", "right"):
        table = doc.add_table(rows=1, cols=2)
        table.autofit = True
        cell_img = table.rows[0].cells[0 if wrap == "left" else 1]
        cell_txt = table.rows[0].cells[1 if wrap == "left" else 0]
        p = cell_img.paragraphs[0]
        run = p.add_run()
        bio.seek(0)
        run.add_picture(bio, width=DocInches(width_in))
        cell_txt.paragraphs[0].add_run("")  # placeholder for surrounding text flow
        return
    p = doc.add_paragraph()
    if wrap == "full":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    bio.seek(0)
    run.add_picture(bio, width=DocInches(width_in))


def _apply_runs(paragraph, runs: list[Any], fallback_text: str = "") -> None:
    """Apply run-level formatting from the editor model."""
    from docx.shared import RGBColor as DocRGB

    if not runs:
        if fallback_text:
            paragraph.add_run(fallback_text)
        return
    for run_spec in runs[:200]:
        if not isinstance(run_spec, dict):
            continue
        text = str(run_spec.get("text") or "")
        if not text:
            continue
        run = paragraph.add_run(text)
        if run_spec.get("bold"):
            run.bold = True
        if run_spec.get("italic"):
            run.italic = True
        if run_spec.get("underline"):
            run.underline = True
        color = run_spec.get("color")
        if isinstance(color, str) and re.match(r"^#[0-9A-Fa-f]{6}$", color):
            h = color[1:]
            run.font.color.rgb = DocRGB(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))
        size = run_spec.get("fontSize") or run_spec.get("font_size")
        try:
            if size:
                run.font.size = Pt(int(size))
        except Exception:
            pass
        font = run_spec.get("font") or run_spec.get("fontFamily")
        if isinstance(font, str) and font.strip():
            run.font.name = font.strip()[:60]


def build_docx_from_blocks(blocks: list[Any], path: Path) -> None:
    """Build a .docx from the editor document model (text + positioned images)."""
    doc = Document()
    for block in blocks[:400]:
        if not isinstance(block, dict):
            continue
        btype = (block.get("type") or "paragraph").lower()
        runs = block.get("runs") if isinstance(block.get("runs"), list) else []
        if btype == "heading":
            level = int(block.get("level") or 1)
            level = 1 if level < 1 else 3 if level > 3 else level
            p = doc.add_heading("", level=level)
            _apply_runs(p, runs, str(block.get("text") or "").strip() or " ")
        elif btype in ("bullet", "list_item", "list"):
            p = doc.add_paragraph(style="List Bullet")
            _apply_runs(p, runs, str(block.get("text") or "").strip())
        elif btype in ("number", "ordered"):
            p = doc.add_paragraph(style="List Number")
            _apply_runs(p, runs, str(block.get("text") or "").strip())
        elif btype == "image":
            src = block.get("src") or block.get("dataUrl") or ""
            blob = _decode_image_src(str(src))
            if not blob:
                continue
            width_in = float(block.get("width_in") or block.get("widthIn") or 4.5)
            wrap = str(block.get("wrap") or "full")
            try:
                _add_picture_wrapped(doc, blob, width_in, wrap)
            except Exception as e:
                print(f"[export] picture skip: {e}")
                cap = doc.add_paragraph(str(block.get("alt") or "[image]"))
                if cap.runs:
                    cap.runs[0].font.size = Pt(10)
        else:
            text = str(block.get("text") or "")
            if text.strip() or runs:
                p = doc.add_paragraph()
                _apply_runs(p, runs, text)
            elif btype == "paragraph":
                doc.add_paragraph("")
    doc.save(str(path))


def _wallet_key(wallet: str) -> str:
    import hashlib

    w = (wallet or "").strip().lower()
    return hashlib.sha256(w.encode("utf-8")).hexdigest()[:40]


def _pass_path(wallet: str) -> Path:
    return _PASS_DIR / f"{_wallet_key(wallet)}.json"


def _load_pass(wallet: str) -> dict[str, Any]:
    p = _pass_path(wallet)
    if not p.is_file():
        return {"free_used": 0, "pass_expires": 0, "updated": 0}
    try:
        data = json.loads(p.read_text(encoding="utf-8"))
        # never keep raw wallet in file payload
        data.pop("wallet", None)
        data.pop("walletAddress", None)
        return data
    except Exception:
        return {"free_used": 0, "pass_expires": 0, "updated": 0}


def _save_pass(wallet: str, data: dict[str, Any]) -> None:
    data = dict(data)
    data["updated"] = time.time()
    data.pop("wallet", None)
    data.pop("walletAddress", None)
    _pass_path(wallet).write_text(json.dumps(data), encoding="utf-8")


def _rate_limit_ok(key: str, limit: int = 40, window_s: int = 600) -> bool:
    k = _wallet_key(key or "anon")
    now = time.time()
    with _rate_lock:
        hits = [t for t in _rate_hits.get(k, []) if now - t < window_s]
        if len(hits) >= limit:
            _rate_hits[k] = hits
            return False
        hits.append(now)
        _rate_hits[k] = hits
        return True


def _fetch_usd_price(cache: dict[str, Any], coingecko_id: str, dex_q: str, fallback: float) -> float:
    import urllib.request as urllib_req

    now = time.time()
    if now - float(cache.get("ts") or 0) < 300 and float(cache.get("price") or 0) > 0:
        return float(cache["price"])
    try:
        req = urllib_req.Request(
            f"https://api.coingecko.com/api/v3/simple/price?ids={coingecko_id}&vs_currencies=usd",
            headers={"User-Agent": "Lightdocs/1.0"},
        )
        with urllib_req.urlopen(req, timeout=8) as r:
            data = json.loads(r.read())
            price = (data.get(coingecko_id) or {}).get("usd")
            if price and float(price) > 0:
                cache["price"] = float(price)
                cache["ts"] = now
                return float(price)
    except Exception:
        pass
    try:
        req = urllib_req.Request(
            f"https://api.dexscreener.com/latest/dex/search?q={dex_q}",
            headers={"User-Agent": "Lightdocs/1.0"},
        )
        with urllib_req.urlopen(req, timeout=8) as r:
            data = json.loads(r.read())
            for pair in data.get("pairs") or []:
                price = float(pair.get("priceUsd") or 0)
                if price > 0:
                    cache["price"] = price
                    cache["ts"] = now
                    return price
    except Exception:
        pass
    cache["ts"] = now
    if float(cache.get("price") or 0) <= 0:
        cache["price"] = fallback
    return float(cache["price"])


def get_lcai_price_usd() -> float:
    return _fetch_usd_price(_lcai_price_cache, "lightchain-ai", "LCAI", 0.004)


def get_keiko_price_usd() -> float:
    # KEIKO may not be on CoinGecko; dex search + env override
    env_p = os.environ.get("KEIKO_USD_PRICE")
    if env_p:
        try:
            return float(env_p)
        except Exception:
            pass
    return _fetch_usd_price(_keiko_price_cache, "keiko", "KEIKO", 0.00001026)


def subscription_amounts() -> dict[str, Any]:
    lcai_usd = get_lcai_price_usd()
    keiko_usd = get_keiko_price_usd()
    lcai_amt = (SUB_PRICE_USD / lcai_usd) if lcai_usd > 0 else 0
    keiko_usd_effective = SUB_PRICE_USD * (1.0 - KEIKO_DISCOUNT)
    keiko_amt = (keiko_usd_effective / keiko_usd) if keiko_usd > 0 else 0
    return {
        "sub_price_usd": SUB_PRICE_USD,
        "keiko_discount": KEIKO_DISCOUNT,
        "keiko_usd_effective": keiko_usd_effective,
        "lcai_price_usd": lcai_usd,
        "keiko_price_usd": keiko_usd,
        "lcai_amount": round(lcai_amt, 6),
        "keiko_amount": round(keiko_amt, 2),
        "pass_days": PASS_DAYS,
        "free_docs_total": FREE_DOCS_TOTAL,
    }


def _authorize_job(wallet: str, device_id: str = "") -> tuple[bool, str, dict[str, Any]]:
    """Monthly pass or one-time free docs per wallet. No per-job credits."""
    w = (wallet or "").strip().lower()
    if not w or not w.startswith("0x") or len(w) < 10:
        return (
            False,
            "Connect your wallet to generate — free beta docs and subscriptions are per wallet.",
            {"need_wallet": True},
        )
    if not _rate_limit_ok(w):
        return False, "Too many requests — wait a few minutes and try again.", {}

    data = _load_pass(w)
    now = time.time()
    expires = float(data.get("pass_expires") or 0)
    if expires > now:
        return True, "subscribed", {
            "paid": True,
            "path": "pass",
            "pass_expires": expires,
            "free_used": int(data.get("free_used") or 0),
            "free_left": max(0, FREE_DOCS_TOTAL - int(data.get("free_used") or 0)),
        }

    free_used = int(data.get("free_used") or 0)
    if free_used < FREE_DOCS_TOTAL:
        data["free_used"] = free_used + 1
        _save_pass(w, data)
        return True, "free", {
            "paid": False,
            "path": "free",
            "free_used": data["free_used"],
            "free_left": FREE_DOCS_TOTAL - data["free_used"],
        }

    return (
        False,
        "Subscribe or renew to keep going — $1/mo (KEIKO at 20% off). Beta pricing may change.",
        {
            "need_pay": True,
            "free_used": free_used,
            "free_left": 0,
            "pass_expires": expires,
        },
    )


@app.get("/api/pay/config")
def pay_config():
    amts = subscription_amounts()
    return jsonify(
        {
            "model": "subscription",
            "sub_price_usd": amts["sub_price_usd"],
            "keiko_discount": amts["keiko_discount"],
            "lcai_amount": amts["lcai_amount"],
            "keiko_amount": amts["keiko_amount"],
            "lcai_price_usd": amts["lcai_price_usd"],
            "keiko_price_usd": amts["keiko_price_usd"],
            "pass_days": PASS_DAYS,
            "free_docs_total": FREE_DOCS_TOTAL,
            "keiko_enabled": bool(KEIKO_RECEIVE_WALLET),
            "lcai_enabled": bool(LCAI_RECEIVE_WALLET),
            "keiko_receive": KEIKO_RECEIVE_WALLET or None,
            "lcai_receive": LCAI_RECEIVE_WALLET or None,
            "keiko_token": "0x93ed20e33e7c88cfa73348086ed1f2c7a2b50854",
            "chain_id": 9200,
            "beta": True,
            "disclaimer": (
                "Beta: Lightdocs is provided as-is. Pricing and terms may change to regular "
                "pricing after testing. No refunds. Service may change or end. AI output is "
                "not professional advice. Wallet connection stays on your device."
            ),
            "privacy": "We store only free-docs-used and pass-expiry per wallet — not your documents.",
        }
    )


@app.get("/api/pay/status")
def pay_status():
    wallet = (request.args.get("wallet") or request.args.get("walletAddress") or "").strip().lower()
    if not wallet:
        return jsonify({"error": "wallet required"}), 400
    data = _load_pass(wallet)
    now = time.time()
    expires = float(data.get("pass_expires") or 0)
    free_used = int(data.get("free_used") or 0)
    return jsonify(
        {
            "model": "subscription",
            "free_used": free_used,
            "free_left": max(0, FREE_DOCS_TOTAL - free_used),
            "free_docs_total": FREE_DOCS_TOTAL,
            "pass_active": expires > now,
            "pass_expires": expires if expires > now else 0,
            "need_subscribe": expires <= now and free_used >= FREE_DOCS_TOTAL,
        }
    )


@app.post("/api/pay/verify-keiko")
def pay_verify_keiko():
    """Verify on-chain KEIKO payment for a 30-day access pass ($1 equivalent at 20% off)."""
    body = request.get_json(silent=True) or {}
    tx_hash = str(body.get("txHash") or body.get("tx_hash") or "").strip()
    wallet = str(body.get("walletAddress") or body.get("wallet") or "").strip().lower()
    if not wallet or not tx_hash:
        return jsonify({"error": "wallet and txHash required"}), 400
    if not KEIKO_RECEIVE_WALLET:
        return jsonify({"error": "KEIKO payments are not configured on this server"}), 503
    amts = subscription_amounts()
    need = float(amts["keiko_amount"] or 0)
    if need <= 0:
        return jsonify({"error": "Could not price KEIKO right now — try again shortly"}), 503
    try:
        import keiko_pay

        used = keiko_pay.UsedTxStore(str(DATA_DIR / "keiko_used_tx.json"))
        ok, err = keiko_pay.register_keiko_payment(
            tx_hash,
            wallet,
            to_wallet=KEIKO_RECEIVE_WALLET,
            amount_keiko=need,
            used_tx_store=used,
        )
        if not ok:
            return jsonify({"error": err or "Payment could not be verified"}), 400
    except Exception as e:
        print(f"[pay] keiko verify failed: {e}")
        return jsonify({"error": "Payment verification failed — try again"}), 400

    data = _load_pass(wallet)
    now = time.time()
    base = max(now, float(data.get("pass_expires") or 0))
    data["pass_expires"] = base + PASS_DAYS * 86400
    _save_pass(wallet, data)
    return jsonify(
        {
            "ok": True,
            "paid_with": "KEIKO",
            "pass_expires": data["pass_expires"],
            "pass_days": PASS_DAYS,
            "amount_keiko": need,
        }
    )


@app.post("/api/export/docx")
def export_docx():
    """Rebuild a .docx from the edited document model (client-side editor source of truth)."""
    body = request.get_json(silent=True) or {}
    blocks = body.get("blocks")
    if not isinstance(blocks, list) or not blocks:
        return jsonify({"error": "Provide blocks[] document model"}), 400
    fname = f"lightdocs-edit-{uuid.uuid4().hex[:8]}.docx"
    path = DATA_DIR / fname
    try:
        build_docx_from_blocks(blocks, path)
    except Exception as e:
        print(f"[export] failed: {e}")
        return jsonify({"error": f"Export failed: {e}"}), 500
    # short-lived: register a mini job meta so cleanup can find it
    job_id = uuid.uuid4().hex
    job = {
        "id": job_id,
        "status": "done",
        "step": "Edited export ready",
        "created": time.time(),
        "file_path": str(path),
        "docx_path": str(path),
        "download_name": fname,
        "mime": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "output": "docx",
        "text_out": None,
        "error": None,
    }
    with _jobs_lock:
        _jobs[job_id] = job
        _save_job(job_id, job)
    return jsonify(
        {
            "job_id": job_id,
            "download_url": f"/api/jobs/{job_id}/download",
            "download_name": fname,
            "output": "docx",
        }
    )


if __name__ == "__main__":
    port = int(os.environ.get("PORT", "8099"))
    app.run(host="0.0.0.0", port=port, debug=False)
